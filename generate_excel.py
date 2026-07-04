import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter
import os

def create_excel_report(summary_csv='rdo_report_v5.csv', mse_csv='frame_mse_report.csv', output_excel='compression_report_v4.xlsx'):
    if not os.path.exists(summary_csv) or not os.path.exists(mse_csv):
        print("CSV files not found.")
        return

    try:
        df_sum = pd.read_csv(summary_csv)
        df_mse = pd.read_csv(mse_csv)
    except Exception as e:
        print(f"Error reading CSV: {e}")
        return

    if df_sum.empty or df_mse.empty:
        print("One of the CSV files is empty.")
        return

    grouped = df_mse.groupby(['Video', 'Config'])
    max_mses = []
    
    for (video, config), group in grouped:
        mses = group['MSE'].values
        max_mse = mses.max() if len(mses) > 0 else 0
        
        if len(mses) >= 5:
            block_mse = max([sum(mses[i:i+5])/5 for i in range(len(mses)-4)])
        else:
            block_mse = max_mse
            
        max_mses.append({'Video': video, 'Config': config, 'MaxMSE': max_mse, 'NoticeableBlockMSE': block_mse})
        
    df_computed = pd.DataFrame(max_mses)
    
    if not df_computed.empty:
        df_final = pd.merge(df_sum, df_computed, on=['Video', 'Config'], how='left')
    else:
        df_final = df_sum
        df_final['MaxMSE'] = 0
        df_final['NoticeableBlockMSE'] = 0

    cols = ['Video', 'Config', 'Size_KB', 'Ratio', 'MSE', 'MaxMSE', 'NoticeableBlockMSE', 'Enc_FPS', 'Dec_FPS']
    df_final = df_final[[c for c in cols if c in df_final.columns]]
    df_final = df_final.rename(columns={'MSE': 'AvgMSE'})

    df_final.to_excel(output_excel, index=False, engine='openpyxl')
    
    wb = openpyxl.load_workbook(output_excel)
    ws = wb.active
    
    red_green = ColorScaleRule(start_type='min', start_color='00FF00', end_type='max', end_color='FF0000')
    green_red = ColorScaleRule(start_type='min', start_color='FF0000', end_type='max', end_color='00FF00')

    cols_map = {col.value: idx for idx, col in enumerate(ws[1], 1)}
    
    for col_name in ['Size_KB', 'Ratio', 'AvgMSE', 'MaxMSE', 'NoticeableBlockMSE']:
        if col_name in cols_map:
            col_letter = get_column_letter(cols_map[col_name])
            ws.conditional_formatting.add(f'{col_letter}2:{col_letter}{ws.max_row}', red_green)
            
    for col_name in ['Enc_FPS', 'Dec_FPS']:
        if col_name in cols_map:
            col_letter = get_column_letter(cols_map[col_name])
            ws.conditional_formatting.add(f'{col_letter}2:{col_letter}{ws.max_row}', green_red)
            
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(cell.value)
            except:
                pass
        ws.column_dimensions[column].width = max_length + 2

    wb.save(output_excel)
    print(f"Report generated: {output_excel}")

if __name__ == '__main__':
    create_excel_report()
